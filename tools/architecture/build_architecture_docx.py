import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


workspace = Path(__file__).resolve().parents[2]
image_path = Path(os.environ.get("AQONE_ARCHITECTURE_PNG", workspace / ".artifacts_build" / "aqone-flowchart" / "aqone-aggregated-architecture.png"))
output_path = workspace / "artifacts" / "architecture" / "AqOne_Aggregated_Technical_Architecture.docx"

document = Document()
section = document.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Inches(22)
section.page_height = Inches(14.7)
section.top_margin = Inches(0.22)
section.bottom_margin = Inches(0.22)
section.left_margin = Inches(0.25)
section.right_margin = Inches(0.25)
section.header_distance = Inches(0)
section.footer_distance = Inches(0)

normal = document.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(9)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1

paragraph = document.add_paragraph()
paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph.paragraph_format.space_before = Pt(0)
paragraph.paragraph_format.space_after = Pt(0)
run = paragraph.add_run()
inline_shape = run.add_picture(str(image_path), height=Inches(14.0))
doc_pr = inline_shape._inline.docPr
doc_pr.set("title", "AqOne Aggregated Technical Architecture")
doc_pr.set(
    "descr",
    "Single-page architecture showing independent data flows for emergency SOS, environmental warnings, trip anomaly detection, drift and search decision support, and consented catch activity.",
)

properties = document.core_properties
properties.title = "AqOne Aggregated Technical Architecture"
properties.subject = "Single-page technical architecture and data-flow diagram"
properties.author = "AqOne"
properties.keywords = "AqOne, architecture, SOS, LoRa, FastAPI, PostgreSQL, flowchart"

output_path.parent.mkdir(parents=True, exist_ok=True)
document.save(output_path)
print(output_path)
