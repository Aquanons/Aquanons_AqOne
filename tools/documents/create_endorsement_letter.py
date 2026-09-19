from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


workspace = Path(__file__).resolve().parents[2]
output = workspace / "artifacts" / "competitions" / "enactus" / "AqOne_Enactus_Endorsement_Letter.docx"
output.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.05


def apply_font(run, size=11):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)


def add_paragraph(text="", *, bold=False, italic=False, align=None, after=6, before=0, size=11):
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    apply_font(run, size)
    return paragraph


title = add_paragraph("Endorsement of HEI Head", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=16, size=14)
title.style = doc.styles["Title"]
title.runs[0].font.color.rgb = None

add_paragraph("[Date to be filled in]", after=16)
add_paragraph("KRISTONI G. GO", bold=True, after=0)
add_paragraph("COO & Country Director", after=0)
add_paragraph("Enactus Philippines", after=14)
add_paragraph("Dear Ms. Go:", after=10)
add_paragraph("Greetings!", after=10)

add_paragraph(
    "This is in reference to the official invitation from Enactus Philippines dated September 11, 2026 entitled \"Official Selection and Invitation to the Enactus Philippines National Competition and Innovation Summit 2026 - Early-Stage Project Competition Track.\"",
    after=8,
)
add_paragraph(
    "We are pleased to submit our institution's official representatives to the Enactus Philippines National Competition and Innovation Summit 2026 - Early-Stage Project Competition Track:",
    after=6,
)

members = [
    ("Faculty Advisor", "Edward S. Gumban"),
    ("Student Member", "Lenard Angelo A. Olajay"),
    ("Student Member", "Daniel Joseph R. Orlina"),
    ("Student Member", "Doreen Kay P. Lachica"),
]
for index, (role, name) in enumerate(members):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.05
    number = paragraph.add_run(f"{index + 1}. ")
    number.bold = True
    apply_font(number)
    text = paragraph.add_run(f"{name} - {role}")
    apply_font(text)

add_paragraph("Project: AqOne", bold=True, after=10, before=4)
add_paragraph(
    "We certify the official participation of the above-named delegation in the competition and affirm the accuracy of the information submitted. The Faculty Advisor is currently employed by Aklan State University - Kalibo Campus, and the three Student Members are currently enrolled for Academic Year 2026-2027. The delegation is authorized to represent the University in the Enactus Philippines National Competition and Innovation Summit 2026.",
    after=8,
)
add_paragraph(
    "For any questions, you may contact Edward S. Gumban through [E-mail Address] and [Contact Number].",
    after=14,
)
add_paragraph("Thank you very much.", after=16)
add_paragraph("Sincerely yours,", after=28)
add_paragraph("[Name and Signature of Registrar Head]", bold=True, after=0)
add_paragraph("[Designation]", after=12)
add_paragraph("Approved by:", bold=True, after=22)
add_paragraph("JEFFREY A. CLARIN, DIT", bold=True, after=0)
add_paragraph("SUC President III", after=0)
add_paragraph("Aklan State University - Kalibo Campus", after=0)

doc.core_properties.title = "Endorsement of HEI Head for Enactus Philippines National Competition 2026"
doc.core_properties.subject = "Endorsement of the AqOne delegation"
doc.core_properties.author = "Aklan State University - Kalibo Campus"
doc.save(output)
print(output.resolve())
